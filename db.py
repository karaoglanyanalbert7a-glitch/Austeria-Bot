"""База данных и игровая логика (без Telegram). Всё, что делает ученик, пишется в таблицу events."""
import hashlib
import json
import random
from collections import Counter, defaultdict
import re
import sqlite3
from datetime import date, datetime, timedelta, timezone

from openpyxl import load_workbook

# ---- Настройки игры: можно менять ----
BASE_XP = 25          # чем больше, тем медленнее растут уровни
MAX_LEVEL = 100
MAX_FREEZES = 2       # максимум «заморозок» серии
CEFR = ["A1", "A2", "B1", "B2", "C1"]
SKILL_STATS = ("vocabulary", "grammar", "listening", "speaking")
XP = {"daily": (10, 2), "practice": (3, 1), "homework": (10, 2), "homework_late": (3, 1)}  # (верный, неверный)
COINS_PER_CORRECT = {"daily": 1, "practice": 0, "homework": 1, "homework_late": 0}
HW_QUIZ_BONUS_XP = 10   # бонус за тест-домашку, сданную в срок
HW_MANUAL_XP = 20       # награда за принятое авторское задание
HW_MANUAL_COINS = 5

SCHEMA = """
CREATE TABLE IF NOT EXISTS groups (
    id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT UNIQUE NOT NULL, target_cefr TEXT, created_at TEXT);
CREATE TABLE IF NOT EXISTS users (
    id INTEGER PRIMARY KEY AUTOINCREMENT, telegram_id INTEGER UNIQUE NOT NULL, nickname TEXT NOT NULL,
    role TEXT NOT NULL DEFAULT 'student', group_id INTEGER REFERENCES groups(id),
    is_active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL, remind INTEGER NOT NULL DEFAULT 1);
CREATE UNIQUE INDEX IF NOT EXISTS ux_users_nick ON users(nickname COLLATE NOCASE);
CREATE TABLE IF NOT EXISTS characters (
    user_id INTEGER PRIMARY KEY REFERENCES users(id),
    xp INTEGER NOT NULL DEFAULT 0, level INTEGER NOT NULL DEFAULT 1, coins INTEGER NOT NULL DEFAULT 0,
    vocabulary INTEGER NOT NULL DEFAULT 0, grammar INTEGER NOT NULL DEFAULT 0,
    listening INTEGER NOT NULL DEFAULT 0, speaking INTEGER NOT NULL DEFAULT 0,
    streak INTEGER NOT NULL DEFAULT 0, freezes INTEGER NOT NULL DEFAULT 1, last_active_date TEXT, band TEXT);
CREATE TABLE IF NOT EXISTS questions (
    id INTEGER PRIMARY KEY, topic TEXT NOT NULL, skill TEXT NOT NULL, cefr TEXT NOT NULL,
    type TEXT NOT NULL, question TEXT NOT NULL,
    option_a TEXT, option_b TEXT, option_c TEXT, option_d TEXT,
    correct TEXT NOT NULL, explanation TEXT, active INTEGER NOT NULL DEFAULT 1);
CREATE TABLE IF NOT EXISTS events (
    id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL REFERENCES users(id),
    ts TEXT NOT NULL, day TEXT NOT NULL, event_type TEXT NOT NULL,
    question_id INTEGER, topic TEXT, skill TEXT, cefr TEXT, is_correct INTEGER,
    response_time_ms INTEGER, context TEXT, extra TEXT, answer_given TEXT, slot TEXT);
CREATE INDEX IF NOT EXISTS ix_events_user ON events(user_id, ts);
CREATE INDEX IF NOT EXISTS ix_events_day ON events(day);
CREATE INDEX IF NOT EXISTS ix_events_uq ON events(user_id, event_type, question_id);
CREATE TABLE IF NOT EXISTS rewards (
    id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL REFERENCES users(id),
    kind TEXT NOT NULL, amount INTEGER NOT NULL, reason TEXT,
    given_by INTEGER REFERENCES users(id), ts TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS homework (
    id INTEGER PRIMARY KEY AUTOINCREMENT, kind TEXT NOT NULL, title TEXT NOT NULL, description TEXT,
    topic TEXT, group_id INTEGER REFERENCES groups(id), due_date TEXT, active INTEGER NOT NULL DEFAULT 1,
    created_by INTEGER REFERENCES users(id), created_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS homework_questions (
    homework_id INTEGER NOT NULL REFERENCES homework(id), question_id INTEGER NOT NULL REFERENCES questions(id),
    pos INTEGER NOT NULL, PRIMARY KEY (homework_id, question_id));
CREATE TABLE IF NOT EXISTS homework_answers (
    homework_id INTEGER NOT NULL, user_id INTEGER NOT NULL, question_id INTEGER NOT NULL,
    is_correct INTEGER NOT NULL, ts TEXT NOT NULL, PRIMARY KEY (homework_id, user_id, question_id));
CREATE TABLE IF NOT EXISTS homework_status (
    homework_id INTEGER NOT NULL, user_id INTEGER NOT NULL, status TEXT NOT NULL, score INTEGER, total INTEGER,
    late INTEGER NOT NULL DEFAULT 0, submitted_at TEXT, completed_at TEXT, teacher_comment TEXT,
    sub_chat INTEGER, sub_msg INTEGER,
    PRIMARY KEY (homework_id, user_id));
CREATE TABLE IF NOT EXISTS topic_info (
    topic TEXT PRIMARY KEY, skill TEXT, level_from TEXT NOT NULL, level_to TEXT NOT NULL,
    ord INTEGER NOT NULL DEFAULT 0, src TEXT NOT NULL DEFAULT 'auto');
CREATE TABLE IF NOT EXISTS student_topics (
    user_id INTEGER NOT NULL REFERENCES users(id), topic TEXT NOT NULL, status TEXT NOT NULL,
    marked_at TEXT NOT NULL, source TEXT, reason TEXT, PRIMARY KEY (user_id, topic));
CREATE TABLE IF NOT EXISTS job_runs (
    job TEXT NOT NULL, key TEXT NOT NULL, ts TEXT NOT NULL, PRIMARY KEY (job, key));
CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS cheatsheets (
    id INTEGER PRIMARY KEY AUTOINCREMENT, topic TEXT UNIQUE NOT NULL, cefr TEXT, text TEXT NOT NULL);
"""


def connect(path):
    con = sqlite3.connect(path)
    con.row_factory = sqlite3.Row
    con.executescript(SCHEMA)
    cols = {r[1] for r in con.execute("PRAGMA table_info(homework_status)")}   # миграция старых баз
    for name in ("sub_chat", "sub_msg"):
        if name not in cols:
            con.execute(f"ALTER TABLE homework_status ADD COLUMN {name} INTEGER")
    if "answer_given" not in {r[1] for r in con.execute("PRAGMA table_info(events)")}:
        con.execute("ALTER TABLE events ADD COLUMN answer_given TEXT")
    if "slot" not in {r[1] for r in con.execute("PRAGMA table_info(events)")}:
        con.execute("ALTER TABLE events ADD COLUMN slot TEXT")
    if "band" not in {r[1] for r in con.execute("PRAGMA table_info(characters)")}:
        con.execute("ALTER TABLE characters ADD COLUMN band TEXT")
    for r in con.execute("SELECT user_id, level FROM characters WHERE band IS NULL").fetchall():
        con.execute("UPDATE characters SET band=? WHERE user_id=?", (cefr_for_level(r["level"]), r["user_id"]))
    if "remind" not in {r[1] for r in con.execute("PRAGMA table_info(users)")}:
        con.execute("ALTER TABLE users ADD COLUMN remind INTEGER NOT NULL DEFAULT 1")
    con.commit()
    return con


def utc_now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


# ---------- уровни ----------
def xp_for_level(n):
    """Сколько всего XP нужно, чтобы достичь уровня n. Рост нелинейный."""
    return int(BASE_XP * (n - 1) ** 1.5)


def level_from_xp(xp):
    lvl = 1
    while lvl < MAX_LEVEL and xp >= xp_for_level(lvl + 1):
        lvl += 1
    return lvl


def cefr_for_level(level):
    return CEFR[min((level - 1) // 20, len(CEFR) - 1)]


def band_start_level(cefr):
    return CEFR.index(cefr) * 20 + 1


def band_top_level(cefr):
    return min(CEFR.index(cefr) * 20 + 20, MAX_LEVEL)


def level_cap_on(con):
    return get_setting(con, "level_cap", "1") == "1"


def capped_level(con, ch, xp):
    """Уровень по опыту, но не выше верха текущей полосы CEFR (пока учитель не переведёт ученика)."""
    lvl = level_from_xp(xp)
    if level_cap_on(con) and ch["band"]:
        lvl = min(lvl, band_top_level(ch["band"]))
    return lvl


# ---------- пользователи ----------
def get_user_by_tg(con, tg_id):
    return con.execute("SELECT * FROM users WHERE telegram_id=?", (tg_id,)).fetchone()


def get_user_by_nick(con, nick):
    """Ищет ник без учёта регистра, в том числе для русских букв."""
    row = con.execute("SELECT * FROM users WHERE nickname=? COLLATE NOCASE", (nick,)).fetchone()
    if row:
        return row
    key = nick.casefold()
    for r in con.execute("SELECT * FROM users"):
        if r["nickname"].casefold() == key:
            return r
    return None


def create_user(con, tg_id, nickname, role="student"):
    cur = con.execute(
        "INSERT INTO users (telegram_id, nickname, role, created_at) VALUES (?,?,?,?)",
        (tg_id, nickname, role, utc_now()))
    if role == "student":  # у учителя нет персонажа и прогресса
        con.execute("INSERT INTO characters (user_id, band) VALUES (?, 'A1')", (cur.lastrowid,))
    con.commit()
    return cur.lastrowid


def drop_teacher_characters(con):
    """Убирает персонажей у учителей (остались от прежних версий)."""
    con.execute("DELETE FROM characters WHERE user_id IN (SELECT id FROM users WHERE role='teacher')")
    con.commit()


def get_character(con, user_id):
    return con.execute("SELECT * FROM characters WHERE user_id=?", (user_id,)).fetchone()


# ---------- журнал событий ----------
def log_event(con, user_id, event_type, day, **kw):
    allowed = ("question_id", "topic", "skill", "cefr", "is_correct", "response_time_ms", "context", "extra",
               "answer_given", "slot")
    data = {k: v for k, v in kw.items() if k in allowed}
    cols = ["user_id", "ts", "day", "event_type"] + list(data)
    vals = [user_id, utc_now(), day.isoformat(), event_type] + list(data.values())
    con.execute(f"INSERT INTO events ({','.join(cols)}) VALUES ({','.join('?' * len(cols))})", vals)
    con.commit()


# ---------- вопросы ----------
def get_question(con, qid):
    return con.execute("SELECT * FROM questions WHERE id=?", (qid,)).fetchone()


def pick_questions(con, user_id, level, n=5):
    """Вопросы уровня ученика и следующего; сначала те, что он давно не видел (или не видел вовсе)."""
    idx = (level - 1) // 20
    levels = {CEFR[idx]}
    if idx + 1 < len(CEFR):
        levels.add(CEFR[idx + 1])
    seen = {r[0]: r[1] for r in con.execute(
        "SELECT question_id, MAX(ts) FROM events WHERE user_id=? AND event_type='answer' AND question_id IS NOT NULL "
        "GROUP BY question_id", (user_id,))}

    def fetch(where, params):
        rows = [dict(r) for r in con.execute("SELECT * FROM questions q WHERE q.active=1 " + where, params)]
        for r in rows:
            r["last_seen"] = seen.get(r["id"])
        return rows

    rows = fetch(f"AND q.cefr IN ({','.join('?' * len(levels))})", sorted(levels))
    if len(rows) < n:
        rows = fetch("", [])
    random.shuffle(rows)
    rows.sort(key=lambda r: r["last_seen"] or "")
    return rows[:n]


def _norm(s):
    s = str(s).replace("’", "'").replace("‘", "'").replace("`", "'")
    return re.sub(r"\s+", " ", s).strip().lower().rstrip(".!?")


def check_answer(q, given):
    if q["type"] == "choice":
        return given.strip().lower() == q["correct"].strip().lower()
    accepted = [_norm(a) for a in q["correct"].split("/")]  # «isn't/is not» = два допустимых ответа
    return _norm(given) in accepted


def correct_display(q):
    if q["type"] == "choice":
        letter = q["correct"].strip().lower()
        return f"{letter.upper()}) {q['option_' + letter]}"
    return " / ".join(a.strip() for a in q["correct"].split("/"))


def _validate(r):
    try:
        int(float(r.get("id", "")))
    except ValueError:
        return "нет номера id"
    for f in ("topic", "skill", "cefr", "type", "question", "correct"):
        if not r.get(f):
            return f"пустое поле «{f}»"
    if r["cefr"].upper() not in CEFR:
        return f"неизвестный уровень «{r['cefr']}»"
    if r["type"].lower() not in ("choice", "text"):
        return f"тип должен быть choice или text, а не «{r['type']}»"
    if r["type"].lower() == "choice":
        opts = [l for l in "abcd" if r.get(f"option_{l}")]
        if len(opts) < 2:
            return "меньше двух вариантов ответа"
        if r["correct"].lower() not in opts:
            return f"правильный ответ «{r['correct']}» не совпадает с вариантами"
    return None


def import_questions(con, path):
    """Читает лист questions из xlsx. Возвращает (сколько загружено, список проблем)."""
    wb = load_workbook(path, data_only=True, read_only=True)
    if "questions" in wb.sheetnames:
        ws = wb["questions"]
    elif "cheatsheets" in wb.sheetnames:
        wb.close()
        return 0, []
    else:
        ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip().lower() if h is not None else "" for h in next(rows)]
    ok, problems = 0, []
    for line, row in enumerate(rows, start=2):
        r = {h: ("" if v is None else str(v).strip()) for h, v in zip(header, row)}
        if not any(r.values()):
            continue
        err = _validate(r)
        if err:
            problems.append(f"строка {line}: {err}")
            continue
        qtype = r["type"].lower()
        active = 0 if r.get("active", "1").lower() in ("0", "0.0", "false", "no", "нет") else 1
        con.execute(
            """INSERT INTO questions (id, topic, skill, cefr, type, question, option_a, option_b, option_c,
                   option_d, correct, explanation, active) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
               ON CONFLICT(id) DO UPDATE SET topic=excluded.topic, skill=excluded.skill, cefr=excluded.cefr,
                   type=excluded.type, question=excluded.question, option_a=excluded.option_a,
                   option_b=excluded.option_b, option_c=excluded.option_c, option_d=excluded.option_d,
                   correct=excluded.correct, explanation=excluded.explanation, active=excluded.active""",
            (int(float(r["id"])), r["topic"], r["skill"].lower(), r["cefr"].upper(), qtype, r["question"],
             r.get("option_a", ""), r.get("option_b", ""), r.get("option_c", ""), r.get("option_d", ""),
             r["correct"].lower() if qtype == "choice" else r["correct"], r.get("explanation", ""), active))
        ok += 1
    con.commit()
    wb.close()
    return ok, problems


# ---------- награды, серия ----------
def apply_answer(con, user_id, is_correct, context, skill):
    xp_ok, xp_bad = XP[context]
    gain = xp_ok if is_correct else xp_bad
    coins = COINS_PER_CORRECT[context] if is_correct else 0
    ch = get_character(con, user_id)
    new_xp = ch["xp"] + gain
    sql, params = "UPDATE characters SET xp=?, level=?, coins=coins+?", [new_xp, capped_level(con, ch, new_xp), coins]
    if is_correct and skill in SKILL_STATS:  # имя колонки проверено по белому списку
        sql += f", {skill}={skill}+1"
    con.execute(sql + " WHERE user_id=?", params + [user_id])
    con.commit()
    return gain, coins


def daily_started_today(con, user_id, today):
    return con.execute(
        "SELECT 1 FROM events WHERE user_id=? AND day=? AND event_type='quiz_start' AND context='daily'",
        (user_id, today.isoformat())).fetchone() is not None


def update_streak(con, user_id, today):
    """Возвращает (серия, потрачена_заморозка, получена_заморозка)."""
    ch = get_character(con, user_id)
    last = date.fromisoformat(ch["last_active_date"]) if ch["last_active_date"] else None
    streak, freezes, used, earned = ch["streak"], ch["freezes"], False, False
    if last == today:
        return streak, False, False
    if last is None:
        streak = 1
    else:
        gap = (today - last).days
        if gap == 1:
            streak += 1
        elif gap == 2 and freezes > 0:
            freezes -= 1
            streak += 1
            used = True
        else:
            streak = 1
    if streak % 7 == 0 and freezes < MAX_FREEZES:
        freezes += 1
        earned = True
    con.execute("UPDATE characters SET streak=?, freezes=?, last_active_date=? WHERE user_id=?",
                (streak, freezes, today.isoformat(), user_id))
    con.commit()
    return streak, used, earned


# ---------- памятки ----------
def import_cheatsheets(con, path):
    """Читает лист cheatsheets (topic, cefr, text). Возвращает (сколько загружено, проблемы)."""
    wb = load_workbook(path, data_only=True, read_only=True)
    if "cheatsheets" not in wb.sheetnames:
        wb.close()
        return 0, []
    rows = wb["cheatsheets"].iter_rows(values_only=True)
    header = [str(h).strip().lower() if h is not None else "" for h in next(rows)]
    ok, problems = 0, []
    for line, row in enumerate(rows, start=2):
        r = {h: ("" if v is None else str(v).strip()) for h, v in zip(header, row)}
        if not any(r.values()):
            continue
        if not r.get("topic") or not r.get("text"):
            problems.append(f"памятки, строка {line}: нужны topic и text")
            continue
        con.execute("""INSERT INTO cheatsheets (topic, cefr, text) VALUES (?,?,?)
                       ON CONFLICT(topic) DO UPDATE SET cefr=excluded.cefr, text=excluded.text""",
                    (r["topic"], r.get("cefr", ""), r["text"]))
        ok += 1
    con.commit()
    wb.close()
    return ok, problems


def import_workbook(con, path):
    """Загружает и вопросы, и памятки. Возвращает (вопросов, проблемы_вопросов, памяток, проблемы_памяток)."""
    q_ok, q_prob = import_questions(con, path)
    c_ok, c_prob = import_cheatsheets(con, path)
    import_topic_info(con, path)     # уровни и порядок тем из листа topics (если он есть)
    derive_topic_info(con)           # для остальных тем уровни определяются по вопросам
    return q_ok, q_prob, c_ok, c_prob


def list_cheatsheets(con):
    return con.execute("SELECT * FROM cheatsheets ORDER BY id").fetchall()


def get_cheatsheet(con, sheet_id):
    return con.execute("SELECT * FROM cheatsheets WHERE id=?", (sheet_id,)).fetchone()


def get_cheatsheet_by_topic(con, topic):
    return con.execute("SELECT * FROM cheatsheets WHERE topic=?", (topic,)).fetchone()


# ---------- группы ----------
def create_group(con, name, target_cefr=None):
    cur = con.execute("INSERT INTO groups (name, target_cefr, created_at) VALUES (?,?,?)",
                      (name, target_cefr, utc_now()))
    con.commit()
    return cur.lastrowid


def get_group(con, name):
    """Ищет группу по названию без учёта регистра, в том числе для русских букв."""
    row = con.execute("SELECT * FROM groups WHERE name=? COLLATE NOCASE", (name,)).fetchone()
    if row:
        return row
    key = name.casefold()
    for r in con.execute("SELECT * FROM groups"):
        if r["name"].casefold() == key:
            return r
    return None


def group_name(con, group_id):
    if not group_id:
        return None
    row = con.execute("SELECT name FROM groups WHERE id=?", (group_id,)).fetchone()
    return row["name"] if row else None


def groups_overview(con):
    return con.execute(
        """SELECT g.id, g.name, g.target_cefr, COUNT(u.id) AS members FROM groups g
           LEFT JOIN users u ON u.group_id=g.id AND u.role='student' AND u.is_active=1
           GROUP BY g.id ORDER BY g.name COLLATE NOCASE""").fetchall()


def set_group(con, user_id, group_id):
    con.execute("UPDATE users SET group_id=? WHERE id=?", (group_id, user_id))
    con.commit()


def delete_group(con, group_id):
    """Удаляет группу, только если в ней никого нет."""
    if con.execute("SELECT 1 FROM users WHERE group_id=?", (group_id,)).fetchone():
        return False
    con.execute("DELETE FROM groups WHERE id=?", (group_id,))
    con.commit()
    return True


def group_students(con, group_id):
    return con.execute("SELECT * FROM users WHERE group_id=? AND role='student' AND is_active=1 "
                       "ORDER BY nickname COLLATE NOCASE", (group_id,)).fetchall()


def group_topic_accuracy(con, group_id, since_day):
    return con.execute(
        """SELECT e.topic AS k, COUNT(*) AS n, SUM(e.is_correct) AS ok FROM events e
           JOIN users u ON u.id=e.user_id
           WHERE u.group_id=? AND e.event_type='answer' AND e.day>=? AND COALESCE(e.slot,'')<>'new'
           GROUP BY e.topic
           ORDER BY 1.0*SUM(e.is_correct)/COUNT(*) ASC""", (group_id, since_day)).fetchall()


# ---------- награды и штрафы ----------
REWARD_KINDS = ("xp", "coins", "freeze")


def adjust(con, user_id, kind, amount, reason, given_by, day):
    """Начисляет (amount > 0) или списывает (amount < 0). Баланс не уходит ниже нуля.
    Возвращает реально применённое изменение (0, если ничего не изменилось)."""
    ch = get_character(con, user_id)
    if kind == "xp":
        new = max(0, ch["xp"] + amount)
        applied = new - ch["xp"]
        con.execute("UPDATE characters SET xp=?, level=? WHERE user_id=?", (new, capped_level(con, ch, new), user_id))
    elif kind == "coins":
        new = max(0, ch["coins"] + amount)
        applied = new - ch["coins"]
        con.execute("UPDATE characters SET coins=? WHERE user_id=?", (new, user_id))
    elif kind == "freeze":
        new = min(MAX_FREEZES, max(0, ch["freezes"] + amount))
        applied = new - ch["freezes"]
        con.execute("UPDATE characters SET freezes=? WHERE user_id=?", (new, user_id))
    else:
        raise ValueError(kind)
    if applied == 0:
        con.commit()
        return 0
    con.execute("INSERT INTO rewards (user_id, kind, amount, reason, given_by, ts) VALUES (?,?,?,?,?,?)",
                (user_id, kind, applied, reason, given_by, utc_now()))
    con.commit()
    log_event(con, user_id, "reward" if applied > 0 else "penalty", day,
              extra=json.dumps({"kind": kind, "amount": applied, "reason": reason}, ensure_ascii=False))
    return applied


# ---------- аналитика для учителя ----------
def students_overview(con, since_day, group_id=None, user_id=None):
    sql = """SELECT u.id, u.nickname, u.group_id, g.name AS grp, c.level, c.streak, c.freezes, c.last_active_date,
             (SELECT MAX(e.day) FROM events e WHERE e.user_id=u.id AND e.event_type='answer') AS last_answer,
             (SELECT COUNT(*) FROM events e WHERE e.user_id=u.id AND e.event_type='answer' AND e.day>=?
                 AND COALESCE(e.slot,'')<>'new') AS n,
             (SELECT COALESCE(SUM(is_correct),0) FROM events e WHERE e.user_id=u.id
                 AND e.event_type='answer' AND e.day>=? AND COALESCE(e.slot,'')<>'new') AS ok
           FROM users u JOIN characters c ON c.user_id=u.id LEFT JOIN groups g ON g.id=u.group_id
           WHERE u.role='student' AND u.is_active=1"""
    params = [since_day, since_day]
    if group_id is not None:
        sql += " AND u.group_id=?"
        params.append(group_id)
    if user_id is not None:
        sql += " AND u.id=?"
        params.append(user_id)
    return con.execute(sql + " ORDER BY u.nickname COLLATE NOCASE", params).fetchall()


def memo_views(con, user_id, since_day):
    return con.execute("SELECT topic AS k, COUNT(*) AS n FROM events WHERE user_id=? "
                       "AND event_type='cheatsheet_view' AND day>=? GROUP BY topic ORDER BY n DESC",
                       (user_id, since_day)).fetchall()


def accuracy_by(con, user_id, column):
    assert column in ("topic", "skill", "cefr")
    return con.execute(
        f"""SELECT {column} AS k, COUNT(*) AS n, SUM(is_correct) AS ok FROM events
            WHERE user_id=? AND event_type='answer' AND COALESCE(slot,'')<>'new' GROUP BY {column}
            ORDER BY 1.0*SUM(is_correct)/COUNT(*) ASC, n DESC""", (user_id,)).fetchall()


def active_days(con, user_id, since_day):
    return con.execute("SELECT COUNT(DISTINCT day) FROM events WHERE user_id=? AND day>=?",
                       (user_id, since_day)).fetchone()[0]


# ---------- домашние задания ----------
def list_topics(con):
    return [r[0] for r in con.execute("SELECT DISTINCT topic FROM questions WHERE active=1 ORDER BY topic")]


def topic_key(topic):
    """Короткий стабильный ключ темы для кнопок (названия тем слишком длинные для callback_data)."""
    return hashlib.md5(topic.encode("utf-8")).hexdigest()[:8]


def topic_by_key(con, key):
    for (t,) in con.execute("SELECT DISTINCT topic FROM questions WHERE active=1"):
        if topic_key(t) == key:
            return t
    return None


def topic_stats(con, cefr=None, query=None):
    """[(тема, число активных вопросов)] с фильтром по уровню и по части названия (без учёта регистра)."""
    sql, params = "SELECT topic, COUNT(*) AS n FROM questions WHERE active=1", []
    if cefr:
        sql += " AND cefr=?"
        params.append(cefr)
    rows = [(r["topic"], r["n"]) for r in con.execute(sql + " GROUP BY topic", params)]
    if query:
        key = query.casefold()
        rows = [r for r in rows if key in r[0].casefold()]
    return sorted(rows, key=lambda r: r[0].casefold())


def count_questions(con, topic=None, cefr=None):
    sql, params = "SELECT COUNT(*) FROM questions WHERE active=1", []
    if topic:
        sql += " AND topic=?"
        params.append(topic)
    if cefr:
        sql += " AND cefr=?"
        params.append(cefr)
    return con.execute(sql, params).fetchone()[0]


def similar_topics(con):
    """Названия тем, которые отличаются только регистром или пробелами: скорее всего, опечатка в таблице."""
    groups = {}
    for (t,) in con.execute("SELECT DISTINCT topic FROM questions"):
        groups.setdefault(" ".join(t.casefold().split()), []).append(t)
    return [sorted(v) for v in groups.values() if len(v) > 1]


def pick_homework_questions(con, topic, count, cefr=None):
    """Случайные вопросы по теме (и уровню), от лёгких к сложным."""
    sql, params = "SELECT id, cefr FROM questions WHERE active=1", []
    if topic:
        sql += " AND topic=?"
        params.append(topic)
    if cefr:
        sql += " AND cefr=?"
        params.append(cefr)
    rows = [dict(r) for r in con.execute(sql, params)]
    random.shuffle(rows)
    rows = rows[:count]
    rows.sort(key=lambda r: CEFR.index(r["cefr"]) if r["cefr"] in CEFR else 9)
    return [r["id"] for r in rows]


def create_homework(con, kind, title, description, topic, group_id, due_date, created_by, question_ids=()):
    cur = con.execute(
        "INSERT INTO homework (kind, title, description, topic, group_id, due_date, created_by, created_at) "
        "VALUES (?,?,?,?,?,?,?,?)", (kind, title, description, topic, group_id, due_date, created_by, utc_now()))
    for pos, qid in enumerate(question_ids):
        con.execute("INSERT INTO homework_questions (homework_id, question_id, pos) VALUES (?,?,?)",
                    (cur.lastrowid, qid, pos))
    con.commit()
    return cur.lastrowid


def get_homework(con, hw_id):
    return con.execute("SELECT * FROM homework WHERE id=?", (hw_id,)).fetchone()


def hw_visible(hw, user):
    return bool(hw and hw["active"] and (hw["group_id"] is None or hw["group_id"] == user["group_id"]))


def homework_for_user(con, user):
    return con.execute(
        """SELECT h.*, s.status, s.score, s.total, s.late, s.teacher_comment FROM homework h
           LEFT JOIN homework_status s ON s.homework_id=h.id AND s.user_id=?
           WHERE h.active=1 AND (h.group_id IS NULL OR h.group_id=?)
           ORDER BY (h.due_date IS NULL), h.due_date, h.id DESC LIMIT 15""",
        (user["id"], user["group_id"])).fetchall()


def hw_question_ids(con, hw_id):
    return [r[0] for r in con.execute(
        "SELECT question_id FROM homework_questions WHERE homework_id=? ORDER BY pos", (hw_id,))]


def hw_remaining(con, hw_id, user_id):
    done = {r[0] for r in con.execute(
        "SELECT question_id FROM homework_answers WHERE homework_id=? AND user_id=?", (hw_id, user_id))}
    return [q for q in hw_question_ids(con, hw_id) if q not in done]


def hw_record_answer(con, hw_id, user_id, question_id, is_correct):
    cur = con.execute("INSERT OR IGNORE INTO homework_answers VALUES (?,?,?,?,?)",
                      (hw_id, user_id, question_id, int(is_correct), utc_now()))
    con.commit()
    return cur.rowcount == 1


def hw_answers_summary(con, hw_id, user_id):
    r = con.execute("SELECT COUNT(*), COALESCE(SUM(is_correct),0) FROM homework_answers "
                    "WHERE homework_id=? AND user_id=?", (hw_id, user_id)).fetchone()
    return r[1], r[0]  # (верных, всего отвечено)


def hw_get_status(con, hw_id, user_id):
    return con.execute("SELECT * FROM homework_status WHERE homework_id=? AND user_id=?",
                       (hw_id, user_id)).fetchone()


def hw_set_status(con, hw_id, user_id, status, **fields):
    allowed = ("score", "total", "late", "submitted_at", "completed_at", "teacher_comment", "sub_chat", "sub_msg")
    keys = [k for k in fields if k in allowed]
    cols = ["homework_id", "user_id", "status"] + keys
    vals = [hw_id, user_id, status] + [fields[k] for k in keys]
    updates = ",".join(f"{c}=excluded.{c}" for c in ["status"] + keys)
    con.execute(f"INSERT INTO homework_status ({','.join(cols)}) VALUES ({','.join('?' * len(cols))}) "
                f"ON CONFLICT(homework_id, user_id) DO UPDATE SET {updates}", vals)
    con.commit()


def hw_reset(con, hw_id, user_id):
    con.execute("DELETE FROM homework_answers WHERE homework_id=? AND user_id=?", (hw_id, user_id))
    con.execute("DELETE FROM homework_status WHERE homework_id=? AND user_id=?", (hw_id, user_id))
    con.commit()


def hw_close(con, hw_id):
    con.execute("UPDATE homework SET active=0 WHERE id=?", (hw_id,))
    con.commit()


def students_in_scope(con, group_id=None):
    sql, params = "SELECT * FROM users WHERE role='student' AND is_active=1", []
    if group_id is not None:
        sql += " AND group_id=?"
        params.append(group_id)
    return con.execute(sql + " ORDER BY nickname COLLATE NOCASE", params).fetchall()


def homework_active(con):
    return con.execute("SELECT h.*, g.name AS grp FROM homework h LEFT JOIN groups g ON g.id=h.group_id "
                       "WHERE h.active=1 ORDER BY h.id DESC LIMIT 20").fetchall()


def hw_progress(con, hw):
    sql = """SELECT u.id, u.nickname, s.status, s.score, s.total, s.late FROM users u
             LEFT JOIN homework_status s ON s.homework_id=? AND s.user_id=u.id
             WHERE u.role='student' AND u.is_active=1"""
    params = [hw["id"]]
    if hw["group_id"] is not None:
        sql += " AND u.group_id=?"
        params.append(hw["group_id"])
    return con.execute(sql + " ORDER BY u.nickname COLLATE NOCASE", params).fetchall()


# ---------- сводки ----------
def daily_status(con, user_id, day):
    """Квиз дня: 'done' (закончен), 'started' (начат, но не закончен) или 'none'."""
    if con.execute("SELECT 1 FROM events WHERE user_id=? AND day=? AND event_type='quiz_complete' "
                   "AND context='daily'", (user_id, day.isoformat())).fetchone():
        return "done"
    return "started" if daily_started_today(con, user_id, day) else "none"


def quiz_counts(con, user_id, since_day):
    """(начато квизов, закончено квизов) за период. Домашние тесты считаются отдельно."""
    def count(event_type):
        return con.execute("SELECT COUNT(*) FROM events WHERE user_id=? AND event_type=? AND day>=?",
                           (user_id, event_type, since_day)).fetchone()[0]
    return count("quiz_start"), count("quiz_complete")


def teacher_overview(con, day):
    one = lambda sql, *p: con.execute(sql, p).fetchone()[0]
    return {
        "students": one("SELECT COUNT(*) FROM users WHERE role='student' AND is_active=1"),
        "groups": one("SELECT COUNT(*) FROM groups"),
        "active_today": one("SELECT COUNT(DISTINCT user_id) FROM events WHERE day=? AND event_type='answer'",
                            day.isoformat()),
        "homeworks": one("SELECT COUNT(*) FROM homework WHERE active=1"),
        "pending": one("SELECT COUNT(*) FROM homework_status s JOIN homework h ON h.id=s.homework_id "
                       "WHERE s.status='submitted' AND h.active=1"),
    }


# ---------- для кнопочного интерфейса учителя ----------
def students_for_picker(con, group_id):
    """Ученики, которых можно добавить в группу: сначала без группы, потом из других групп."""
    return con.execute(
        """SELECT u.id, u.nickname, u.group_id, g.name AS grp FROM users u LEFT JOIN groups g ON g.id=u.group_id
           WHERE u.role='student' AND u.is_active=1 AND (u.group_id IS NULL OR u.group_id<>?)
           ORDER BY (u.group_id IS NOT NULL), g.name COLLATE NOCASE, u.nickname COLLATE NOCASE""",
        (group_id,)).fetchall()


def hw_pending(con):
    """Сданные авторские работы, которые ждут проверки."""
    return con.execute(
        """SELECT s.homework_id, s.user_id, u.nickname, h.title, s.sub_chat, s.sub_msg
           FROM homework_status s JOIN homework h ON h.id=s.homework_id JOIN users u ON u.id=s.user_id
           WHERE s.status='submitted' AND h.active=1 ORDER BY s.submitted_at""").fetchall()


def set_group_level(con, group_id, cefr):
    con.execute("UPDATE groups SET target_cefr=? WHERE id=?", (cefr, group_id))
    con.commit()


def rename_group(con, group_id, name):
    con.execute("UPDATE groups SET name=? WHERE id=?", (name, group_id))
    con.commit()


def make_teacher(con, user_id):
    """Превращает ошибочно зарегистрированного «ученика» в учителя (если забыли TEACHER_IDS при первом запуске)."""
    con.execute("UPDATE users SET role='teacher', group_id=NULL WHERE id=?", (user_id,))
    con.execute("DELETE FROM characters WHERE user_id=?", (user_id,))
    con.commit()


# ---------- фоновые задачи: чтобы после перезапуска ничего не отправлялось повторно ----------
def job_done(con, job, key):
    return con.execute("SELECT 1 FROM job_runs WHERE job=? AND key=?", (job, key)).fetchone() is not None


def job_mark(con, job, key):
    con.execute("INSERT OR IGNORE INTO job_runs (job, key, ts) VALUES (?,?,?)", (job, key, utc_now()))
    con.commit()


def get_setting(con, key, default="1"):
    row = con.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
    return row["value"] if row else default


def set_setting(con, key, value):
    con.execute("INSERT INTO settings (key, value) VALUES (?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                (key, str(value)))
    con.commit()


def set_remind(con, user_id, flag):
    con.execute("UPDATE users SET remind=? WHERE id=?", (1 if flag else 0, user_id))
    con.commit()


def effective_streak(ch, today):
    """Серия, которую ученик реально имеет сейчас. Сохранённое число обновляется только после квиза,
    поэтому у пропавшего ученика оно устаревает: считаем сгоревшую серию нулём."""
    if not ch["last_active_date"]:
        return 0
    gap = (today - date.fromisoformat(ch["last_active_date"])).days
    if gap <= 1 or (gap == 2 and ch["freezes"] > 0):
        return ch["streak"]
    return 0


# =====================================================================================
#  Программа: уровни тем, пройденные темы, перевод на уровень, сборка квиза
# =====================================================================================
LEVEL_WEIGHTS = {1: 0.20, 0: 0.55, -1: 0.18, -2: 0.07}   # смещение уровня вопроса относительно уровня ученика
RELEARN_WINDOW = 10      # тема «пройдена», но на повтор уходит, если среди последних 10 ответов…
RELEARN_MIN = 6          # …(их не меньше 6)…
RELEARN_BELOW = 50       # …верных меньше 50%
COVER_TARGET = 0.85      # доля пройденных тем полосы, после которой бот предлагает перевод на следующий уровень
MIN_LEVEL_QUESTIONS = 30 # меньше стольких вопросов уровня в банке: предупреждаем учителя


def parse_cefr_range(text):
    """«A1-A2» → ('A1', 'A2'); «A2» → ('A2', 'A2'); непонятное → None."""
    parts = [p for p in re.split(r"[-–/,; ]+", (text or "").upper()) if p in CEFR]
    if not parts:
        return None
    return min(parts, key=CEFR.index), max(parts, key=CEFR.index)


def import_topic_info(con, path):
    """Лист topics: topic, typical_cefr (или cefr/level), необязательно skill и order. Порядок строк = порядок в программе."""
    wb = load_workbook(path, data_only=True, read_only=True)
    if "topics" not in wb.sheetnames:
        wb.close()
        return 0
    rows = wb["topics"].iter_rows(values_only=True)
    header = [str(h).strip().lower() if h is not None else "" for h in next(rows)]
    n = 0
    for i, row in enumerate(rows, start=1):
        r = {h: ("" if v is None else str(v).strip()) for h, v in zip(header, row)}
        rng = parse_cefr_range(r.get("typical_cefr") or r.get("cefr") or r.get("level") or "")
        if not r.get("topic") or not rng:
            continue
        try:
            order = int(float(r.get("order") or i))
        except ValueError:
            order = i
        con.execute(
            "INSERT INTO topic_info (topic, skill, level_from, level_to, ord, src) VALUES (?,?,?,?,?, 'sheet') "
            "ON CONFLICT(topic) DO UPDATE SET skill=COALESCE(excluded.skill, topic_info.skill), "
            "level_from=excluded.level_from, level_to=excluded.level_to, ord=excluded.ord, src='sheet'",
            (r["topic"], (r.get("skill") or "").lower() or None, rng[0], rng[1], order))
        n += 1
    con.commit()
    wb.close()
    return n


def derive_topic_info(con):
    """Навык и уровни тем, которых нет в листе topics, определяются по вопросам:
    навык = самый частый у темы, уровни = те, где в теме не меньше 3 вопросов."""
    per = defaultdict(lambda: {"skills": Counter(), "levels": Counter(), "min_id": 10 ** 9})
    for r in con.execute("SELECT topic, skill, cefr, id FROM questions WHERE active=1"):
        p = per[r["topic"]]
        p["skills"][r["skill"]] += 1
        p["levels"][r["cefr"]] += 1
        p["min_id"] = min(p["min_id"], r["id"])
    existing = {r["topic"]: r for r in con.execute("SELECT * FROM topic_info")}
    for rank, (topic, p) in enumerate(sorted(per.items(), key=lambda kv: kv[1]["min_id"])):
        skill = p["skills"].most_common(1)[0][0]
        known = [l for l in CEFR if p["levels"][l] >= 3] or [min(p["levels"], key=CEFR.index)]
        ex = existing.get(topic)
        if ex is None:
            con.execute("INSERT INTO topic_info (topic, skill, level_from, level_to, ord, src) VALUES (?,?,?,?,?, 'auto')",
                        (topic, skill, known[0], known[-1], 1000 + rank))
        elif ex["src"] == "auto":
            con.execute("UPDATE topic_info SET skill=?, level_from=?, level_to=? WHERE topic=?",
                        (skill, known[0], known[-1], topic))
        elif not ex["skill"]:
            con.execute("UPDATE topic_info SET skill=? WHERE topic=?", (skill, topic))
    con.commit()


def ensure_topic_info(con):
    if con.execute("SELECT 1 FROM questions q LEFT JOIN topic_info t ON t.topic=q.topic "
                   "WHERE q.active=1 AND t.topic IS NULL LIMIT 1").fetchone():
        derive_topic_info(con)


def grammar_topics(con):
    """Темы грамматики в порядке программы: [{topic, level_from, level_to, ord}]."""
    ensure_topic_info(con)
    return [dict(r) for r in con.execute(
        "SELECT * FROM topic_info WHERE skill='grammar' AND topic IN (SELECT DISTINCT topic FROM questions "
        "WHERE active=1) ORDER BY ord, topic")]


def level_question_counts(con):
    return Counter({r[0]: r[1] for r in con.execute("SELECT cefr, COUNT(*) FROM questions WHERE active=1 GROUP BY cefr")})


# ---------- отметки «пройдено» ----------
def covered_topics(con, user_id):
    return {r["topic"]: dict(r) for r in con.execute("SELECT * FROM student_topics WHERE user_id=?", (user_id,))}


def set_topic_status(con, user_id, topic, status, source, reason=None):
    con.execute("INSERT INTO student_topics (user_id, topic, status, marked_at, source, reason) VALUES (?,?,?,?,?,?) "
                "ON CONFLICT(user_id, topic) DO UPDATE SET status=excluded.status, marked_at=excluded.marked_at, "
                "source=excluded.source, reason=excluded.reason",
                (user_id, topic, status, utc_now(), source, reason))
    con.commit()


def clear_topic(con, user_id, topic):
    con.execute("DELETE FROM student_topics WHERE user_id=? AND topic=?", (user_id, topic))
    con.commit()


def program_summary(con, user_id):
    ch = get_character(con, user_id)
    band = ch["band"] or cefr_for_level(ch["level"])
    bi = CEFR.index(band)
    topics = grammar_topics(con)
    cov = covered_topics(con, user_id)
    prog = [t for t in topics if CEFR.index(t["level_from"]) <= bi <= CEFR.index(t["level_to"])]
    done = [t for t in prog if cov.get(t["topic"], {}).get("status") == "covered"]
    return {"band": band, "program": len(prog), "done": len(done),
            "covered_total": sum(1 for t in topics if cov.get(t["topic"], {}).get("status") == "covered"),
            "grammar_total": len(topics),
            "relearn": [dict(t, reason=cov[t["topic"]]["reason"]) for t in topics
                        if cov.get(t["topic"], {}).get("status") == "relearn"],
            "next_band": CEFR[bi + 1] if bi + 1 < len(CEFR) else None,
            "ready": len(prog) >= 5 and len(done) / len(prog) >= COVER_TARGET}


def check_relearn(con, user_id, topic):
    """Пройденная тема уходит «на повтор», если после отметки ученик стабильно ошибается в ней."""
    row = con.execute("SELECT * FROM student_topics WHERE user_id=? AND topic=? AND status='covered'",
                      (user_id, topic)).fetchone()
    if not row:
        return None
    rows = con.execute("SELECT is_correct FROM events WHERE user_id=? AND event_type='answer' AND topic=? AND ts>=? "
                       "AND COALESCE(slot,'')<>'new' ORDER BY id DESC LIMIT ?",
                       (user_id, topic, row["marked_at"], RELEARN_WINDOW)).fetchall()
    n, ok = len(rows), sum(r[0] or 0 for r in rows)
    if n >= RELEARN_MIN and 100 * ok / n < RELEARN_BELOW:
        reason = f"верно {round(100 * ok / n)}% за последние {n} ответов"
        set_topic_status(con, user_id, topic, "relearn", "auto", reason)
        return {"topic": topic, "reason": reason}
    return None


# ---------- перевод на уровень ----------
def jump_plan(con, user_id, cefr):
    ch = get_character(con, user_id)
    cur = ch["band"] or cefr_for_level(ch["level"])
    idx, cur_idx = CEFR.index(cefr), CEFR.index(cur)
    topics, cov = grammar_topics(con), covered_topics(con, user_id)
    is_cov = lambda t: cov.get(t["topic"], {}).get("status") == "covered"
    to_cover = [t["topic"] for t in topics if CEFR.index(t["level_from"]) < idx and not is_cov(t)]
    to_uncover = ([t["topic"] for t in topics if CEFR.index(t["level_from"]) >= idx and is_cov(t)]
                  if idx < cur_idx else [])
    start = band_start_level(cefr)
    xp = xp_for_level(start)
    if idx >= cur_idx:
        xp = max(xp, ch["xp"])
    level = level_from_xp(xp)
    if level_cap_on(con):
        level = min(level, band_top_level(cefr))
    n_q = level_question_counts(con)[cefr]
    return {"cefr": cefr, "band_old": cur, "level_old": ch["level"], "xp_old": ch["xp"], "level": level, "xp": xp,
            "to_cover": to_cover, "to_uncover": to_uncover, "questions": n_q, "thin": n_q < MIN_LEVEL_QUESTIONS}


def apply_jump(con, user_id, cefr, day, by_user_id=None):
    plan = jump_plan(con, user_id, cefr)
    con.execute("UPDATE characters SET band=?, xp=?, level=? WHERE user_id=?", (cefr, plan["xp"], plan["level"], user_id))
    con.commit()
    for t in plan["to_cover"]:
        set_topic_status(con, user_id, t, "covered", "level_jump")
    for t in plan["to_uncover"]:
        clear_topic(con, user_id, t)
    log_event(con, user_id, "level_set", day,
              extra=json.dumps({"from": plan["band_old"], "to": cefr, "level": plan["level"], "xp": plan["xp"],
                                "covered": len(plan["to_cover"]), "uncovered": len(plan["to_uncover"]),
                                "by": by_user_id}, ensure_ascii=False))
    return plan


# ---------- сборка ежедневного квиза ----------
def compose_quiz(con, user_id, size=8, today=None):
    """Собирает квиз: грамматика по пройденным темам + разведка по новым темам + лексика/чтение по уровням.
    Правила: вопрос никогда не выше уровня ученика +1; уровни вопросов выбираются с весами LEVEL_WEIGHTS
    (в основном свой уровень, иногда на 1 выше и ниже, редко ещё ниже), поэтому ни «потолок» A1, ни «вечное A1»
    у старших. Если у ученика нет ни одной отмеченной темы, квиз собирается только по уровням."""
    today = today or date.today()
    ch = get_character(con, user_id)
    li = CEFR.index(cefr_for_level(ch["level"]))
    seen = {r[0]: r[1] for r in con.execute(
        "SELECT question_id, MAX(ts) FROM events WHERE user_id=? AND event_type='answer' AND question_id IS NOT NULL "
        "GROUP BY question_id", (user_id,))}
    questions = []
    for r in con.execute("SELECT * FROM questions WHERE active=1"):
        q = dict(r)
        if CEFR.index(q["cefr"]) <= li + 1:                 # жёсткий потолок: не выше уровня ученика +1
            q["last_seen"] = seen.get(q["id"])
            questions.append(q)
    cov = covered_topics(con, user_id)
    covered = {t for t, r in cov.items() if r["status"] == "covered"}
    relearn = {t for t, r in cov.items() if r["status"] == "relearn"}
    chosen, used_q, used_t = [], set(), set()

    # насколько каждой пройденной теме нужно повторение: слабая или давно не встречалась
    last_day = {r[0]: r[1] for r in con.execute(
        "SELECT topic, MAX(day) FROM events WHERE user_id=? AND event_type='answer' GROUP BY topic", (user_id,))}
    recent = {r[0]: (r[1], r[2] or 0) for r in con.execute(
        "SELECT topic, COUNT(*), SUM(is_correct) FROM events WHERE user_id=? AND event_type='answer' AND day>=? "
        "AND COALESCE(slot,'')<>'new' GROUP BY topic", (user_id, (today - timedelta(days=30)).isoformat()))}

    def topic_need(topic):
        n, ok = recent.get(topic, (0, 0))
        if n >= 3 and ok / n < 0.6:
            return 3
        ld = last_day.get(topic)
        return 2 if ld is None or (today - date.fromisoformat(ld)).days >= 14 else 1

    def pick(pool, k, slot, by_topic=False):
        for _ in range(k):
            cand = [q for q in pool if q["id"] not in used_q]
            if not cand:
                return
            buckets = defaultdict(list)
            for q in cand:
                buckets[max(CEFR.index(q["cefr"]) - li, -2)].append(q)
            # вес уровня, но корзина из нескольких вопросов выбирается реже, чтобы они не повторялись каждый квиз
            w = {off: LEVEL_WEIGHTS.get(off, 0) * min(1.0, len(buckets[off]) / 10)
                 for off in buckets if LEVEL_WEIGHTS.get(off, 0) > 0} or {o: 1 for o in buckets}
            bucket = buckets[random.choices(list(w), weights=list(w.values()))[0]]
            fresh = [q for q in bucket if q["topic"] not in used_t] or bucket
            if by_topic:      # тему выбираем с учётом потребности в повторении
                topics = sorted({q["topic"] for q in fresh})
                topic = random.choices(topics, weights=[topic_need(t) for t in topics])[0]
                fresh = [q for q in fresh if q["topic"] == topic]
            random.shuffle(fresh)
            q = min(fresh, key=lambda x: x["last_seen"] or "")       # непросмотренные и давно виденные — первыми
            chosen.append(dict(q, slot=slot))
            used_q.add(q["id"])
            used_t.add(q["topic"])

    if not covered and not relearn:                           # программа ещё не заполнена: только по уровням
        pick(questions, size, "level")
    else:
        info = {t["topic"]: t for t in grammar_topics(con)}
        n_new = 0 if size < 3 else (2 if size >= 8 else 1)
        n_level = max(1, round(size * 0.35))
        # разведка: непройденные темы не выше уровня ученика; сначала те, что вернулись на повтор, потом по порядку программы
        new_topics = sorted((t for t in info.values() if t["topic"] not in covered and CEFR.index(t["level_from"]) <= li),
                            key=lambda t: (t["topic"] not in relearn, t["ord"]))
        for t in new_topics[:n_new]:
            pool = [q for q in questions if q["topic"] == t["topic"]]
            if pool:
                q = min(pool, key=lambda x: (x["last_seen"] is not None, CEFR.index(x["cefr"]), x["last_seen"] or ""))
                chosen.append(dict(q, slot="new"))
                used_q.add(q["id"])
                used_t.add(q["topic"])
        n_review = max(0, size - len(chosen) - n_level)
        pick([q for q in questions if q["topic"] in covered], n_review, "review", by_topic=True)
        pick([q for q in questions if q["skill"] != "grammar"], n_level, "level")
    if len(chosen) < size:        # добор: не хватило вопросов в каком-то слоте
        pick(questions, size - len(chosen), "level")
    random.shuffle(chosen)
    if chosen and chosen[0]["slot"] == "new":                 # начинаем не с незнакомой темы
        for i, q in enumerate(chosen):
            if q["slot"] != "new":
                chosen[0], chosen[i] = chosen[i], chosen[0]
                break
    return chosen[:size]
