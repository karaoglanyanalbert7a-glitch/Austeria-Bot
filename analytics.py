"""Аналитика для учителя: флаги риска, динамика по неделям, недельный дайджест, экспорт в Excel.
Модуль работает только с базой данных (без Telegram), поэтому его удобно проверять на тестовых данных."""
from collections import Counter, defaultdict, deque
from datetime import date, timedelta
from html import escape as esc

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import db

# ---- пороги: можно менять ----
ABSENT_DAYS = 5          # не занимается столько дней и больше → флаг «пропал»
NEVER_DAYS = 3           # зарегистрировался давно, но ни разу не занимался
DECLINE_POINTS = 15      # точность упала на столько процентных пунктов и больше
MIN_ANSWERS = 8          # минимум ответов в каждой неделе, чтобы сравнивать точность
OVERDUE_MIN = 2          # просрочено столько заданий и больше → флаг
DROPOUT_MIN_STARTS = 3   # минимум начатых квизов за 14 дней, чтобы судить, что ученик их бросает
WEEKS_DEFAULT = 6


def d(s):
    return date.fromisoformat(s)


def pct(ok, n):
    return round(100 * ok / n) if n else None


def pct_text(ok, n):
    return f"{pct(ok, n)}%" if n else "—"


def monday(day):
    return day - timedelta(days=day.weekday())


def bar(value, width=10):
    """Текстовая полоска 0–100%."""
    if value is None:
        return "░" * width
    filled = max(0, min(width, round(width * value / 100)))
    return "█" * filled + "░" * (width - filled)


def split_text(text, limit=3500):
    """Режет длинный текст по строкам на части, пригодные для сообщений Telegram."""
    chunks, cur = [], ""
    for line in text.split("\n"):
        if len(cur) + len(line) + 1 > limit and cur:
            chunks.append(cur.rstrip("\n"))
            cur = ""
        cur += line + "\n"
    if cur.strip():
        chunks.append(cur.rstrip("\n"))
    return chunks


# =====================================================================================
#  Выборки
# =====================================================================================
def answer_rows(con, since_day, until_day=None, user_id=None, group_id=None):
    """Ответы учеников за период: (user_id, day, is_correct, topic, skill)."""
    sql = ("SELECT e.user_id, e.day, e.is_correct, e.topic, e.skill FROM events e JOIN users u ON u.id=e.user_id "
           "WHERE e.event_type='answer' AND u.role='student' AND u.is_active=1 AND e.day>=?")
    params = [since_day.isoformat()]
    if until_day:
        sql += " AND e.day<=?"
        params.append(until_day.isoformat())
    if user_id is not None:
        sql += " AND e.user_id=?"
        params.append(user_id)
    if group_id is not None:
        sql += " AND u.group_id=?"
        params.append(group_id)
    return con.execute(sql, params).fetchall()


def totals(rows):
    n = len(rows)
    return n, sum(r["is_correct"] or 0 for r in rows)


# =====================================================================================
#  Динамика по неделям
# =====================================================================================
def weekly_stats(con, today, weeks=WEEKS_DEFAULT, user_id=None, group_id=None):
    """Список недель (от старой к новой): начало недели, ответов, верных, число активных учеников."""
    first = monday(today) - timedelta(weeks=weeks - 1)
    buckets = {first + timedelta(weeks=i): {"n": 0, "ok": 0, "users": set()} for i in range(weeks)}
    for r in answer_rows(con, first, user_id=user_id, group_id=group_id):
        b = buckets.get(monday(d(r["day"])))
        if b is not None:
            b["n"] += 1
            b["ok"] += r["is_correct"] or 0
            b["users"].add(r["user_id"])
    return [{"week": w, "n": b["n"], "ok": b["ok"], "active": len(b["users"])} for w, b in sorted(buckets.items())]


def weekly_lines(con, today, weeks=WEEKS_DEFAULT, user_id=None, group_id=None):
    """Текстовая «диаграмма» динамики: по строке на неделю."""
    stats = weekly_stats(con, today, weeks, user_id, group_id)
    while stats and not stats[0]["n"]:      # недели до первого ответа не показываем
        stats.pop(0)
    if not stats:
        return []
    lines = []
    for s in stats:
        label = f"{s['week'].strftime('%d.%m')}"
        acc = pct(s["ok"], s["n"])
        extra = f" · активных: {s['active']}" if group_id is not None else ""
        lines.append(f"{label} {bar(acc)} {pct_text(s['ok'], s['n']):>4} · {s['n']} отв.{extra}")
    return lines


# =====================================================================================
#  Флаги риска
# =====================================================================================
def _iso_week(day):
    y, w, _ = day.isocalendar()
    return f"{y}-{w:02d}"


def risk_flags(con, today, end=None, user_id=None):
    """Возвращает {id ученика: [(код, текст, ключ_эпизода)]} только для тех, у кого есть сигналы.
    Ключ эпизода нужен, чтобы не присылать одно и то же предупреждение каждый день."""
    end = end or today
    cur_start, prev_start, prev_end = end - timedelta(days=6), end - timedelta(days=13), end - timedelta(days=7)
    sql = "SELECT id, nickname, created_at, group_id FROM users WHERE role='student' AND is_active=1"
    params = []
    if user_id is not None:
        sql += " AND id=?"
        params.append(user_id)
    students = con.execute(sql, params).fetchall()
    last = {r[0]: r[1] for r in con.execute(
        "SELECT user_id, MAX(day) FROM events WHERE event_type='answer' GROUP BY user_id")}
    cur = defaultdict(list)
    prev = defaultdict(list)
    for r in answer_rows(con, prev_start, end, user_id=user_id):
        (cur if d(r["day"]) >= cur_start else prev)[r["user_id"]].append(r)
    since14 = (today - timedelta(days=13)).isoformat()
    starts = dict(con.execute("SELECT user_id, COUNT(*) FROM events WHERE event_type='quiz_start' AND day>=? "
                              "GROUP BY user_id", (since14,)).fetchall())
    finishes = dict(con.execute("SELECT user_id, COUNT(*) FROM events WHERE event_type='quiz_complete' AND day>=? "
                                "GROUP BY user_id", (since14,)).fetchall())
    out = {}
    for s in students:
        uid, flags = s["id"], []
        la = last.get(uid)
        if la:
            gap = (today - d(la)).days
            if gap >= ABSENT_DAYS:
                flags.append(("absent", f"😴 не занимается {gap} дн.", f"absent:{la}"))
        elif (today - d(s["created_at"][:10])).days >= NEVER_DAYS:
            flags.append(("never", "😴 зарегистрировался(ась), но ни разу не занимался(ась)", "never"))
        n1, ok1 = totals(cur[uid])
        n0, ok0 = totals(prev[uid])
        if n1 >= MIN_ANSWERS and n0 >= MIN_ANSWERS and pct(ok1, n1) <= pct(ok0, n0) - DECLINE_POINTS:
            flags.append(("decline", f"📉 точность упала с {pct(ok0, n0)}% до {pct(ok1, n1)}%",
                          f"decline:{_iso_week(today)}"))
        overdue = [r[0] for r in con.execute(
            """SELECT h.id FROM homework h LEFT JOIN homework_status st ON st.homework_id=h.id AND st.user_id=?
               WHERE h.active=1 AND (h.group_id IS NULL OR h.group_id=?) AND h.due_date IS NOT NULL AND h.due_date<?
               AND COALESCE(st.status,'new') NOT IN ('done','submitted') ORDER BY h.id""",
            (uid, s["group_id"], today.isoformat()))]
        if len(overdue) >= OVERDUE_MIN:
            flags.append(("overdue", f"⏰ просрочено заданий: {len(overdue)}", "overdue:" + ",".join(map(str, overdue))))
        st_n, fin_n = starts.get(uid, 0), finishes.get(uid, 0)
        if st_n >= DROPOUT_MIN_STARTS and fin_n / st_n <= 0.5:
            flags.append(("dropout", f"🚪 бросает квизы: закончено {fin_n} из {st_n}", f"dropout:{_iso_week(today)}"))
        if flags:
            out[uid] = flags
    return out


def flags_text(flags):
    return "; ".join(f[1] for f in flags)


# =====================================================================================
#  Светофор учеников
# =====================================================================================
STATUS_ICON = {"green": "🟢", "yellow": "🟡", "red": "🔴", "none": "⚪"}
LEGEND = "🟢 всё хорошо · 🟡 присмотреться · 🔴 нужна помощь · ⚪ мало данных"


def student_status(con, today, user_id=None):
    """Один значок на ученика: {id: {"code", "icon", "reason"}}.
    🔴 пропал/ни разу не занимался или сразу несколько проблем, 🟡 одна проблема или низкая точность/редкие занятия,
    ⚪ слишком мало данных (новичок), 🟢 всё в порядке."""
    flags = risk_flags(con, today, user_id=user_id)
    sql, params = "SELECT id, created_at FROM users WHERE role='student' AND is_active=1", []
    if user_id is not None:
        sql += " AND id=?"
        params.append(user_id)
    rows14 = defaultdict(list)
    for r in answer_rows(con, today - timedelta(days=13), user_id=user_id):
        rows14[r["user_id"]].append(r)
    out = {}
    for s in con.execute(sql, params).fetchall():
        uid, fl = s["id"], flags.get(s["id"], [])
        codes = {f[0] for f in fl}
        reg_days = (today - d(s["created_at"][:10])).days
        n14, ok14 = totals(rows14[uid])
        days7 = len({r["day"] for r in rows14[uid] if d(r["day"]) >= today - timedelta(days=6)})
        if codes & {"absent", "never"} or len(codes) >= 2:
            code, reason = "red", flags_text(fl)
        elif codes:
            code, reason = "yellow", flags_text(fl)
        elif n14 >= 10 and pct(ok14, n14) < 50:
            code, reason = "yellow", f"🎯 точность ниже 50% ({pct(ok14, n14)}% за 14 дней)"
        elif reg_days >= 7 and days7 <= 1:
            code, reason = "yellow", "занимается редко"
        elif n14 < 3:
            code, reason = "none", "мало данных"
        else:
            code, reason = "green", "всё в порядке"
        out[uid] = {"code": code, "icon": STATUS_ICON[code], "reason": reason}
    return out


def status_line(statuses):
    c = Counter(v["code"] for v in statuses.values())
    return f"🚦 Светофор: 🟢 {c['green']} · 🟡 {c['yellow']} · 🔴 {c['red']} · ⚪ {c['none']}"


# =====================================================================================
#  Карта освоения тем
# =====================================================================================
MASTERY_WINDOW = 10      # состояние темы считается по последним 10 ответам
MASTERY_ICON = {"mastered": "🟢", "progress": "🟡", "problem": "🔴"}
MASTERY_WORD = {"mastered": "освоено", "progress": "в процессе", "problem": "проблема"}


def mastery_state(n, acc):
    """🔴 проблема: ≥3 ответов и точность <50%. 🟢 освоено: ≥5 ответов и точность ≥80%. Иначе 🟡 в процессе."""
    if n >= 3 and acc < 50:
        return "problem"
    if n >= 5 and acc >= 80:
        return "mastered"
    return "progress"


def topic_mastery(con, user_id=None, group_id=None):
    """{(id ученика, тема): {state, n, ok, acc, total}} по последним ответам темы."""
    sql = ("SELECT e.user_id, e.topic, e.is_correct FROM events e JOIN users u ON u.id=e.user_id "
           "WHERE e.event_type='answer' AND e.topic IS NOT NULL AND u.role='student' AND u.is_active=1")
    params = []
    if user_id is not None:
        sql += " AND e.user_id=?"
        params.append(user_id)
    if group_id is not None:
        sql += " AND u.group_id=?"
        params.append(group_id)
    win, total = defaultdict(lambda: deque(maxlen=MASTERY_WINDOW)), defaultdict(int)
    for r in con.execute(sql + " ORDER BY e.id", params):
        win[(r["user_id"], r["topic"])].append(r["is_correct"] or 0)
        total[(r["user_id"], r["topic"])] += 1
    out = {}
    for k, w in win.items():
        n, ok = len(w), sum(w)
        out[k] = {"state": mastery_state(n, pct(ok, n)), "n": n, "ok": ok, "acc": pct(ok, n), "total": total[k]}
    return out


def student_map(con, user_id):
    """Темы ученика: сначала проблемные, потом «в процессе», потом освоенные."""
    order = {"problem": 0, "progress": 1, "mastered": 2}
    items = [dict(topic=k[1], **v) for k, v in topic_mastery(con, user_id=user_id).items()]
    items.sort(key=lambda x: (order[x["state"]], x["acc"], x["topic"].casefold()))
    return items


def group_map(con, group_id=None):
    """Темы группы (или всех учеников): у скольких учеников проблема / в процессе / освоено."""
    names = {r["id"]: r["nickname"] for r in con.execute("SELECT id, nickname FROM users")}
    by_topic = defaultdict(lambda: {"problem": [], "progress": [], "mastered": []})
    for (uid, topic), v in topic_mastery(con, group_id=group_id).items():
        by_topic[topic][v["state"]].append(names[uid])
    rows = [{"topic": t, "problem": len(v["problem"]), "progress": len(v["progress"]), "mastered": len(v["mastered"]),
             "total": sum(len(x) for x in v.values()), "problem_names": sorted(v["problem"])}
            for t, v in by_topic.items()]
    rows.sort(key=lambda r: (-r["problem"], -r["progress"], r["topic"].casefold()))
    return rows


# =====================================================================================
#  Типичные ошибки
# =====================================================================================
def common_mistakes(con, since_day, group_id=None, user_id=None, limit=8, min_wrong=3):
    """Вопросы, на которых чаще всего ошибаются, и какой неверный ответ выбирают чаще всего."""
    sql = ("SELECT e.question_id, e.is_correct, e.answer_given, q.topic, q.cefr, q.question, q.type, q.correct, "
           "q.option_a, q.option_b, q.option_c, q.option_d FROM events e JOIN questions q ON q.id=e.question_id "
           "JOIN users u ON u.id=e.user_id WHERE e.event_type='answer' AND e.day>=? AND u.role='student' "
           "AND u.is_active=1")
    params = [since_day.isoformat()]
    if group_id is not None:
        sql += " AND u.group_id=?"
        params.append(group_id)
    if user_id is not None:
        sql += " AND e.user_id=?"
        params.append(user_id)
    per = {}
    for r in con.execute(sql, params):
        p = per.setdefault(r["question_id"], {"row": r, "n": 0, "wrong": 0, "given": Counter()})
        p["n"] += 1
        if not r["is_correct"]:
            p["wrong"] += 1
            if r["answer_given"]:
                p["given"][r["answer_given"].strip().lower()] += 1
    items = []
    for p in per.values():
        if p["wrong"] < min_wrong:
            continue
        r, top = p["row"], p["given"].most_common(1)
        items.append({"topic": r["topic"], "cefr": r["cefr"], "question": r["question"], "correct": db.correct_display(r),
                      "n": p["n"], "wrong": p["wrong"], "top": top[0][0] if top else None,
                      "top_count": top[0][1] if top else 0, "recorded": sum(p["given"].values())})
    items.sort(key=lambda x: (-x["wrong"], -x["wrong"] / x["n"]))
    return items[:limit]


def mistakes_text(items, start=1):
    lines = []
    for i, x in enumerate(items, start):
        lines.append(f"{i}. {esc(x['topic'])} · {x['cefr']} — «{esc(x['question'][:90])}»")
        tail = f"   Верно: {esc(x['correct'])}. Ошибок: {x['wrong']} из {x['n']}"
        if x["top"]:
            tail += f". Чаще всего выбирали: «{esc(x['top'])}» ({x['top_count']} из {x['recorded']})"
        lines.append(tail)
    return "\n".join(lines)


# =====================================================================================
#  Недельный дайджест
# =====================================================================================
def _arrow(now, before, unit="", points=False):
    if before is None or now is None:
        return ""
    diff = now - before
    if diff == 0:
        return " (без изменений)"
    sign = "▲ +" if diff > 0 else "▼ −"
    return f" ({sign}{abs(diff)}{' п.п.' if points else unit})"


def build_digest(con, today):
    """Текст дайджеста (HTML) за последние 7 завершённых дней и сравнение с предыдущими 7."""
    end = today - timedelta(days=1)
    cur_start, prev_start, prev_end = end - timedelta(days=6), end - timedelta(days=13), end - timedelta(days=7)
    students = con.execute("SELECT u.id, u.nickname, u.group_id, g.name AS grp FROM users u LEFT JOIN groups g "
                           "ON g.id=u.group_id WHERE u.role='student' AND u.is_active=1 "
                           "ORDER BY u.nickname COLLATE NOCASE").fetchall()
    head = f"📬 <b>Недельный дайджест</b> · {cur_start.strftime('%d.%m')}–{end.strftime('%d.%m')}"
    if not students:
        return head + "\n\nУчеников пока нет."
    cur_rows = answer_rows(con, cur_start, end)
    prev_rows = answer_rows(con, prev_start, prev_end)
    n1, ok1 = totals(cur_rows)
    n0, ok0 = totals(prev_rows)
    act1, act0 = {r["user_id"] for r in cur_rows}, {r["user_id"] for r in prev_rows}
    lines = [head, ""]
    if not n1:
        lines.append("На этой неделе никто не отвечал на вопросы.")
    else:
        lines += ["<b>Итоги недели</b>",
                  f"👥 Занимались: {len(act1)} из {len(students)}{_arrow(len(act1), len(act0) or None)}",
                  f"✍️ Ответов: {n1}{_arrow(n1, n0 or None)}",
                  f"🎯 Точность: {pct_text(ok1, n1)}{_arrow(pct(ok1, n1), pct(ok0, n0), points=True)}"]
    lines += ["", status_line(student_status(con, today))]
    # по группам
    by_group = defaultdict(list)
    for s in students:
        by_group[s["grp"] or "Без группы"].append(s["id"])
    if len(by_group) > 1 or "Без группы" not in by_group:
        lines += ["", "<b>По группам</b>"]
        for name, ids in sorted(by_group.items()):
            rows = [r for r in cur_rows if r["user_id"] in ids]
            n, ok = totals(rows)
            lines.append(f"• {esc(name)} — занимались {len({r['user_id'] for r in rows})}/{len(ids)} · "
                         f"{pct_text(ok, n)} · {n} отв.")
    # темы с проблемами (по карте освоения) и типичные ошибки
    problem = [r for r in group_map(con) if r["problem"] >= 2 or (r["problem"] >= 1 and r["problem"] / r["total"] >= 0.4)][:3]
    if problem:
        lines += ["", "<b>Темы, где больше всего проблем</b>"]
        lines += [f"• {esc(r['topic'])}: 🔴 {r['problem']} · 🟡 {r['progress']} · 🟢 {r['mastered']} (из {r['total']})"
                  for r in problem]
    else:
        topics = defaultdict(lambda: [0, 0])
        for r in cur_rows:
            topics[r["topic"]][0] += 1
            topics[r["topic"]][1] += r["is_correct"] or 0
        weak = [w for w in sorted((pct(ok, n), t, n) for t, (n, ok) in topics.items() if n >= 10)[:3] if w[0] < 75]
        if weak:
            lines += ["", "<b>Что стоит разобрать на уроке</b>"] + [f"• {esc(t)}: {p}% ({n} отв.)" for p, t, n in weak]
    mist = common_mistakes(con, cur_start, limit=3)
    if mist:
        lines += ["", "<b>🧩 Типичные ошибки недели</b>", mistakes_text(mist)]
    # внимание
    flags = risk_flags(con, today, end=end)
    if flags:
        names = {s["id"]: s["nickname"] for s in students}
        lines += ["", "<b>⚠️ Требуют внимания</b>"]
        lines += [f"• {esc(names[uid])} — {esc(flags_text(f))}" for uid, f in list(flags.items())[:12]]
        if len(flags) > 12:
            lines.append(f"…и ещё {len(flags) - 12}")
    # самые активные
    counts = defaultdict(int)
    for r in cur_rows:
        counts[r["user_id"]] += 1
    if counts:
        names = {s["id"]: s["nickname"] for s in students}
        top = sorted(counts.items(), key=lambda kv: -kv[1])[:3]
        lines += ["", "<b>🏆 Самые активные:</b> " + ", ".join(f"{esc(names[u])} ({c})" for u, c in top)]
    # домашка
    hw_lines = []
    for h in db.homework_active(con):
        prog = db.hw_progress(con, h)
        if not prog:
            continue
        done = sum(1 for p in prog if p["status"] == "done")
        waiting = sum(1 for p in prog if p["status"] == "submitted")
        title = esc(h["title"][:40])
        if h["due_date"] and h["due_date"] < today.isoformat():
            missing = [p["nickname"] for p in prog if p["status"] not in ("done", "submitted")]
            if missing:
                hw_lines.append(f"• ⚠️ «{title}» просрочено: не сделали {esc(', '.join(missing[:6]))}"
                                + ("…" if len(missing) > 6 else ""))
                continue
        hw_lines.append(f"• «{title}» ({'до ' + d(h['due_date']).strftime('%d.%m') if h['due_date'] else 'без срока'}): "
                        f"выполнили {done}/{len(prog)}" + (f", ждёт проверки: {waiting}" if waiting else ""))
    if hw_lines:
        lines += ["", "<b>📝 Домашние задания</b>"] + hw_lines[:8]
    # динамика
    dyn = weekly_lines(con, end)          # последняя строка — неделя, закончившаяся вчера
    if dyn:
        lines += ["", "<b>📈 Динамика (точность по неделям)</b>", "<pre>" + "\n".join(dyn) + "</pre>"]
    return "\n".join(lines)


# =====================================================================================
#  Экспорт в Excel
# =====================================================================================
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", start_color="1F4E78")
BODY_FONT = Font(name="Arial", size=10)
GREEN, RED = PatternFill("solid", start_color="C6EFCE"), PatternFill("solid", start_color="FFC7CE")


def _sheet(wb, title, header, rows, widths=None, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    for r in rows:
        ws.append(list(r))
    for c in ws[1]:
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY_FONT
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, max(11, min(34, len(str(h)) + 4)))
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


def build_export(con, today, path, weeks=8, log_days=180):
    """Собирает Excel-файл: ученики, темы по ученикам, динамика, журнал ответов, домашка."""
    since = today - timedelta(days=13)
    flags = risk_flags(con, today)
    status = student_status(con, today)
    students = db.students_overview(con, (today - timedelta(days=6)).isoformat())
    wb = Workbook()

    rows = []
    for s in students:
        uid = s["id"]
        ch = db.get_character(con, uid)
        tot = con.execute("SELECT COUNT(*), COALESCE(SUM(is_correct),0) FROM events WHERE user_id=? "
                          "AND event_type='answer'", (uid,)).fetchone()
        started, finished = db.quiz_counts(con, uid, since.isoformat())
        last = con.execute("SELECT MAX(day) FROM events WHERE user_id=? AND event_type='answer'", (uid,)).fetchone()[0]
        rows.append([s["nickname"], s["grp"] or "", status[uid]["icon"] + " " + status[uid]["reason"], ch["level"], db.cefr_for_level(ch["level"]), ch["xp"], ch["coins"],
                     db.effective_streak(ch, today), last or "", tot[0], pct(tot[1], tot[0]), s["n"],
                     pct(s["ok"], s["n"]), started, finished, flags_text(flags.get(uid, []))])
    ws = _sheet(wb, "Ученики", ["Ник", "Группа", "Светофор", "Уровень", "CEFR", "XP", "Монеты", "Серия, дн.", "Последний ответ",
                                "Ответов всего", "Точность всего, %", "Ответов за 7 дн.", "Точность за 7 дн., %",
                                "Квизов начато (14 дн.)", "Квизов закончено (14 дн.)", "Сигналы"], rows,
                {1: 18, 2: 18, 3: 40, 16: 48}, first=True)

    # темы × ученики
    topics = [r[0] for r in con.execute("SELECT DISTINCT topic FROM events WHERE event_type='answer' AND topic IS NOT NULL "
                                        "ORDER BY topic")]
    matrix = []
    for s in students:
        row = [s["nickname"], s["grp"] or ""]
        stat = {r["k"]: pct(r["ok"], r["n"]) for r in db.accuracy_by(con, s["id"], "topic")}
        row += [stat.get(t) for t in topics]
        matrix.append(row)
    ws = _sheet(wb, "Темы × ученики, %", ["Ник", "Группа"] + topics, matrix, {1: 18, 2: 18})
    for row in ws.iter_rows(min_row=2, min_col=3):
        for c in row:
            if isinstance(c.value, int) and (c.value >= 80 or c.value < 50):
                c.fill = GREEN if c.value >= 80 else RED

    # освоение тем
    mast = topic_mastery(con)
    m_topics = sorted({k[1] for k in mast}, key=str.casefold)
    m_rows = [[s["nickname"], s["grp"] or ""] + [MASTERY_WORD[mast[(s["id"], t)]["state"]] if (s["id"], t) in mast else None
                                                 for t in m_topics] for s in students]
    ws = _sheet(wb, "Освоение тем", ["Ник", "Группа"] + m_topics, m_rows, {1: 18, 2: 18})
    fills = {"освоено": GREEN, "проблема": RED, "в процессе": PatternFill("solid", start_color="FFEB9C")}
    for row in ws.iter_rows(min_row=2, min_col=3):
        for c in row:
            if c.value in fills:
                c.fill = fills[c.value]

    # динамика
    first_week = monday(today) - timedelta(weeks=weeks - 1)
    week_starts = [first_week + timedelta(weeks=i) for i in range(weeks)]
    dyn = []
    for s in students:
        stats = {w["week"]: pct(w["ok"], w["n"]) for w in weekly_stats(con, today, weeks, user_id=s["id"])}
        dyn.append([s["nickname"], s["grp"] or ""] + [stats.get(w) for w in week_starts])
    total = {w["week"]: w for w in weekly_stats(con, today, weeks)}
    dyn.append(["ВСЕ", ""] + [pct(total[w]["ok"], total[w]["n"]) for w in week_starts])
    dyn.append(["Ответов (все)", ""] + [total[w]["n"] for w in week_starts])
    ws = _sheet(wb, "Динамика, %", ["Ник", "Группа"] + [f"нед. {w.strftime('%d.%m')}" for w in week_starts], dyn,
                {1: 18, 2: 18})
    for row in ws.iter_rows(min_row=2, max_row=1 + len(students), min_col=3):
        for c in row:
            if isinstance(c.value, int) and (c.value >= 80 or c.value < 50):
                c.fill = GREEN if c.value >= 80 else RED

    # журнал ответов
    log = con.execute(
        """SELECT e.day, u.nickname, g.name AS grp, e.topic, e.skill, e.cefr, e.question_id, e.is_correct,
                  e.response_time_ms, e.context FROM events e JOIN users u ON u.id=e.user_id
           LEFT JOIN groups g ON g.id=u.group_id WHERE e.event_type='answer' AND u.role='student' AND e.day>=?
           ORDER BY e.day, e.id LIMIT 50000""", ((today - timedelta(days=log_days)).isoformat(),)).fetchall()
    _sheet(wb, "Журнал ответов", ["Дата", "Ник", "Группа", "Тема", "Навык", "CEFR", "id вопроса", "Верно (1/0)",
                                  "Время, мс", "Откуда"], [tuple(r) for r in log], {1: 12, 2: 18, 3: 18, 4: 28})

    # типичные ошибки
    mist = common_mistakes(con, today - timedelta(days=60), limit=100, min_wrong=2)
    _sheet(wb, "Типичные ошибки", ["Тема", "Уровень", "Вопрос", "Верный ответ", "Ответов", "Ошибок", "Частая ошибка",
                                   "Сколько раз"],
           [[x["topic"], x["cefr"], x["question"], x["correct"], x["n"], x["wrong"], x["top"] or "", x["top_count"] or ""]
            for x in mist], {1: 24, 3: 50, 4: 22, 7: 22})

    # домашка
    hw_rows = []
    for h in con.execute("SELECT * FROM homework ORDER BY id").fetchall():
        for p in db.hw_progress(con, h):
            hw_rows.append([h["id"], h["title"], "тест" if h["kind"] == "quiz" else "авторское",
                            db.group_name(con, h["group_id"]) or "все", h["due_date"] or "", "да" if h["active"] else "нет",
                            p["nickname"], p["status"] or "new", p["score"], p["total"], "да" if p["late"] else ""])
    _sheet(wb, "Домашка", ["№", "Задание", "Тип", "Группа", "Срок", "Активно", "Ученик", "Статус", "Балл", "Из",
                           "Опоздание"], hw_rows, {2: 30, 7: 18})
    wb.save(path)
    return path
